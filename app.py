import os
import json
from io import BytesIO
from datetime import datetime, timezone, timedelta
from flask import Flask, render_template, request, jsonify, redirect, url_for, session, send_file
from dotenv import load_dotenv
from supabase import create_client, Client
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "super_secret_attendance_key")

# --- SUPABASE CONFIGURATION ---
SUPABASE_URL = os.environ.get("SUPABASE_URL")
SUPABASE_KEY = os.environ.get("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    raise ValueError("SUPABASE_URL and SUPABASE_KEY environment variables are required.")

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

SCAN_INTERVAL_MINUTES = 15
PH_TZ = timezone(timedelta(hours=8))

# --- NAVIGATION ROUTES ---

@app.route('/')
def home():
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    error = None
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        
        if username == 'admin' and password == 'admin123':
            session['logged_in'] = True
            return redirect(url_for('admin'))
        else:
            error = 'Invalid admin credentials. Please try again.'
            
    return render_template('login.html', error=error)

@app.route('/kiosk')
def kiosk():
    session.pop('logged_in', None)
    return render_template('index.html')

@app.route('/admin')
def admin():
    if not session.get('logged_in'):
        return redirect(url_for('login'))
    return render_template('admin.html')

@app.route('/record')
def record():
    if not session.get('logged_in'):
        return redirect(url_for('login'))
    return render_template('record.html')

@app.route('/logout')
def logout():
    session.pop('logged_in', None)
    return redirect(url_for('login'))

# --- AUTH VERIFICATION API ---

@app.route('/api/verify_admin', methods=['POST'])
def verify_admin():
    data = request.json or {}
    password = data.get('password')
    
    if password == 'admin123':
        session['logged_in'] = True
        return jsonify({"status": "success", "redirect": url_for('admin')})
    else:
        return jsonify({"status": "error", "message": "Incorrect password. Access denied!"}), 401

# --- EMPLOYEE API ---

@app.route('/api/employees', methods=['GET'])
def get_employees():
    res = supabase.table("employees").select("*").execute()
    rows = res.data or []

    employees = []
    for r in rows:
        face_desc = r.get("face_descriptor")
        if isinstance(face_desc, str):
            face_desc = json.loads(face_desc)

        employees.append({
            "emp_code": r["emp_code"],
            "name": r["name"],
            "age": r["age"],
            "gender": r["gender"],
            "birthday": str(r["birthday"]),
            "address": r["address"],
            "contact": r.get("contact", ""),
            "pin": r["pin"],
            "face_descriptor": face_desc
        })
    return jsonify(employees)

@app.route('/api/register', methods=['POST'])
def register_employee():
    data = request.json
    try:
        new_employee = {
            "emp_code": data['emp_code'],
            "name": data['name'],
            "age": data['age'],
            "gender": data['gender'],
            "birthday": data['birthday'],
            "address": data['address'],
            "contact": data.get('contact', ''),
            "pin": str(data['pin']),
            "face_descriptor": data['face_descriptor']
        }
        supabase.table("employees").insert(new_employee).execute()
        return jsonify({"status": "success", "message": f"Employee {data['name']} registered successfully!"})
    except Exception as e:
        err_msg = str(e)
        if "duplicate key" in err_msg or "unique constraint" in err_msg:
            return jsonify({"status": "error", "message": "Employee Code already exists!"}), 400
        return jsonify({"status": "error", "message": err_msg}), 400

@app.route('/api/employees/<emp_code>', methods=['PUT'])
def update_employee(emp_code):
    data = request.json
    update_data = {
        "name": data['name'],
        "age": data['age'],
        "gender": data['gender'],
        "birthday": data['birthday'],
        "address": data['address'],
        "contact": data.get('contact', ''),
        "pin": str(data['pin'])
    }
    if data.get('face_descriptor'):
        update_data["face_descriptor"] = data['face_descriptor']

    supabase.table("employees").update(update_data).eq("emp_code", emp_code).execute()
    return jsonify({"status": "success", "message": "Employee record updated successfully!"})

@app.route('/api/employees/<emp_code>', methods=['DELETE'])
def delete_employee(emp_code):
    supabase.table("employees").delete().eq("emp_code", emp_code).execute()
    return jsonify({"status": "success", "message": f"Employee {emp_code} deleted!"})

# --- AUTO-LOGGING ATTENDANCE API ---

@app.route('/api/clock', methods=['POST'])
def clock():
    data = request.json
    emp_code = data.get('emp_code')
    method = data.get('method')
    pin = data.get('pin', None)

    emp_res = supabase.table("employees").select("name", "pin").eq("emp_code", emp_code).execute()
    if not emp_res.data:
        return jsonify({"status": "error", "message": "Employee ID not recognized!"}), 404

    employee = emp_res.data[0]
    emp_name = employee["name"]
    stored_pin = str(employee["pin"])

    if method == 'MANUAL' and str(pin) != stored_pin:
        return jsonify({"status": "error", "message": "Invalid Security PIN!"}), 401

    now_utc = datetime.now(timezone.utc)
    now_ph = now_utc.astimezone(PH_TZ)

    last_log_res = supabase.table("attendance") \
        .select("timestamp") \
        .eq("emp_code", emp_code) \
        .order("id", desc=True) \
        .limit(1) \
        .execute()

    if last_log_res.data and last_log_res.data[0].get("timestamp"):
        ts_str = last_log_res.data[0]["timestamp"].replace('Z', '+00:00')
        try:
            last_time = datetime.fromisoformat(ts_str)
            if last_time.tzinfo is None:
                last_time = last_time.replace(tzinfo=timezone.utc)
            else:
                last_time = last_time.astimezone(timezone.utc)

            time_diff_seconds = (now_utc - last_time).total_seconds()
            cooldown_seconds = SCAN_INTERVAL_MINUTES * 60

            if 0 <= time_diff_seconds < cooldown_seconds:
                remaining_secs = int(cooldown_seconds - time_diff_seconds)
                rem_mins = remaining_secs // 60
                rem_secs = remaining_secs % 60
                time_str = f"{rem_mins}m {rem_secs}s" if rem_mins > 0 else f"{rem_secs}s"
                return jsonify({
                    "status": "error",
                    "message": f"Scan rejected! {emp_name} must wait {time_str} before scanning again."
                }), 400
        except Exception as e:
            print(f"Error checking cooldown timestamp: {e}")

    today_ph_date = now_ph.strftime("%Y-%m-%d")

    all_user_logs = supabase.table("attendance") \
        .select("type, timestamp") \
        .eq("emp_code", emp_code) \
        .order("id", desc=False) \
        .execute()

    today_types = []
    for row in (all_user_logs.data or []):
        ts = row.get("timestamp")
        if not ts:
            continue
        try:
            dt = datetime.fromisoformat(ts.replace('Z', '+00:00'))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            if dt.astimezone(PH_TZ).strftime("%Y-%m-%d") == today_ph_date:
                today_types.append(row["type"].strip().upper())
        except Exception:
            continue

    if 'IN' not in today_types:
        action_type = 'IN'
    elif 'LUNCH OUT' not in today_types and 'LUNCH_OUT' not in today_types:
        action_type = 'LUNCH OUT'
    elif 'LUNCH IN' not in today_types and 'LUNCH_IN' not in today_types:
        action_type = 'LUNCH IN'
    elif 'OUT' not in today_types:
        action_type = 'OUT'
    else:
        return jsonify({
            "status": "error", 
            "message": f"All daily punches already completed for {emp_name}!"
        }), 400

    supabase.table("attendance").insert({
        "emp_code": emp_code,
        "name": emp_name,
        "timestamp": now_utc.isoformat(),
        "type": action_type,
        "method": method
    }).execute()

    return jsonify({
        "status": "success",
        "message": f"Recorded {action_type} for {emp_name}",
        "action_type": action_type,
        "employee": emp_name
    })

def get_aggregated_logs():
    raw_res = supabase.table("attendance").select("*").order("timestamp", desc=False).execute()
    raw_logs = raw_res.data or []

    daily_records = {}
    for log in raw_logs:
        log_id = log["id"]
        emp_code = log["emp_code"]
        name = log["name"]
        ts_str = log.get("timestamp", "")
        ctype = log.get("type", "")
        method = log.get("method", "")

        date_str, time_str = "", ""
        if ts_str:
            try:
                dt = datetime.fromisoformat(ts_str.replace('Z', '+00:00'))
                if dt.tzinfo is None:
                    dt = dt.replace(tzinfo=timezone.utc)
                dt_ph = dt.astimezone(PH_TZ)
                date_str = dt_ph.strftime("%Y-%m-%d")
                time_str = dt_ph.strftime("%H:%M:%S")
            except Exception:
                pass

        key = (emp_code, date_str)
        if key not in daily_records:
            daily_records[key] = {
                'id': log_id,
                'emp_code': emp_code,
                'name': name,
                'date': date_str,
                'time_in': '-',
                'lunch_out': '-',
                'lunch_in': '-',
                'time_out': '-',
                'method': method
            }

        ctype_clean = ctype.strip().upper() if ctype else ''
        if ctype_clean == 'IN':
            daily_records[key]['time_in'] = time_str
        elif ctype_clean in ['LUNCH OUT', 'LUNCH_OUT']:
            daily_records[key]['lunch_out'] = time_str
        elif ctype_clean in ['LUNCH IN', 'LUNCH_IN']:
            daily_records[key]['lunch_in'] = time_str
        elif ctype_clean == 'OUT':
            daily_records[key]['time_out'] = time_str

    return sorted(list(daily_records.values()), key=lambda x: x['id'], reverse=True)

@app.route('/api/records', methods=['GET'])
def get_records():
    return jsonify(get_aggregated_logs())

@app.route('/api/records/update', methods=['PUT'])
def update_attendance_record():
    if not session.get('logged_in'):
        return jsonify({"status": "error", "message": "Unauthorized"}), 401

    data = request.json or {}
    emp_code = data.get('emp_code')
    date_str = data.get('date')
    
    if not emp_code or not date_str:
        return jsonify({"status": "error", "message": "Missing employee code or date."}), 400

    times = {
        'IN': data.get('time_in'),
        'LUNCH OUT': data.get('lunch_out'),
        'LUNCH IN': data.get('lunch_in'),
        'OUT': data.get('time_out')
    }

    emp_res = supabase.table("employees").select("name").eq("emp_code", emp_code).execute()
    emp_name = emp_res.data[0]["name"] if emp_res.data else "Unknown"

    for ctype, time_val in times.items():
        if not time_val or time_val.strip() in ['-', '']:
            continue
        
        clean_time = time_val.strip()
        time_parts = clean_time.split(':')
        if len(time_parts) == 2:
            clean_time = f"{clean_time}:00"
            
        dt_ph = datetime.fromisoformat(f"{date_str}T{clean_time}").replace(tzinfo=PH_TZ)
        full_ts = dt_ph.astimezone(timezone.utc).isoformat()

        day_start = f"{date_str}T00:00:00"
        day_end = f"{date_str}T23:59:59"

        existing = supabase.table("attendance") \
            .select("id") \
            .eq("emp_code", emp_code) \
            .gte("timestamp", day_start) \
            .lte("timestamp", day_end) \
            .in_("type", [ctype, ctype.replace(' ', '_')]) \
            .execute()

        if existing.data:
            supabase.table("attendance").update({"timestamp": full_ts}).eq("id", existing.data[0]["id"]).execute()
        else:
            supabase.table("attendance").insert({
                "emp_code": emp_code,
                "name": emp_name,
                "timestamp": full_ts,
                "type": ctype,
                "method": 'MANUAL'
            }).execute()

    return jsonify({"status": "success", "message": "Attendance record updated successfully!"})

@app.route('/api/records/delete', methods=['DELETE'])
def delete_attendance_record():
    if not session.get('logged_in'):
        return jsonify({"status": "error", "message": "Unauthorized"}), 401

    data = request.json or {}
    emp_code = data.get('emp_code')
    date_str = data.get('date')

    if not emp_code or not date_str:
        return jsonify({"status": "error", "message": "Missing employee code or date."}), 400

    day_start = f"{date_str}T00:00:00"
    day_end = f"{date_str}T23:59:59"

    try:
        supabase.table("attendance") \
            .delete() \
            .eq("emp_code", emp_code) \
            .gte("timestamp", day_start) \
            .lte("timestamp", day_end) \
            .execute()
        return jsonify({"status": "success", "message": f"Attendance record for {emp_code} on {date_str} deleted successfully!"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/export_excel', methods=['GET'])
def export_excel():
    records = get_aggregated_logs()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"
    ws.views.sheetView[0].showGridLines = True

    HEADER_BG, ZEBRA_BG, WHITE, BORDER_COLOR = "0F172A", "1E293B", "FFFFFF", "334155"

    font_title = Font(name="Segoe UI", size=16, bold=True, color="38BDF8")
    font_header = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
    font_data = Font(name="Segoe UI", size=11, color="E2E8F0")
    
    fill_header = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR), right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR), bottom=Side(style='thin', color=BORDER_COLOR)
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A1:H1")
    ws["A1"] = "MASTER ATTENDANCE REPORT"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_left
    ws.row_dimensions[1].height = 30

    headers = ["Log ID", "Employee Code", "Name", "Time In", "Lunch Out", "Lunch In", "Time Out", "Method Used"]
    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 25

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font, cell.fill, cell.alignment, cell.border = font_header, fill_header, align_center, thin_border

    for row_idx, r in enumerate(records, start=4):
        row_data = [f"#{r['id']}", r['emp_code'], r['name'], r['time_in'], r['lunch_out'], r['lunch_in'], r['time_out'], r['method']]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20

        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font, cell.border = font_data, thin_border
            cell.alignment = align_left if col_idx == 3 else align_center
            if is_even:
                cell.fill = fill_zebra

    for col in ws.columns:
        max_len = max([len(str(cell.value or '')) for cell in col if cell.row >= 3] or [10])
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="Attendance_Report.xlsx")

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port)